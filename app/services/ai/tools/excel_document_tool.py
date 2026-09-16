"""Controlled Excel document tools for configured agents."""
from __future__ import annotations

import json
import re
from collections import Counter
from datetime import date, datetime, time as dt_time
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.cell import range_boundaries

from app.core.context import get_current_agent_context
from app.services.ai.tools.document_paths import (
    DocumentPathError,
    resolve_document_input_path,
    resolve_document_output_path,
)
from app.services.ai.tools.generated_file_service import register_artifact
from app.services.ai.tools.tool_compat import tool

_EXTENSIONS = {".xlsx"}
_MAX_PREVIEW_ROWS = 20
_MAX_PREVIEW_COLUMNS = 20
_MAX_RANGE_ROWS = 1000
_MAX_RANGE_COLUMNS = 50
_MAX_PROFILE_SHEETS = 20
_MAX_PROFILE_COLUMNS = 20
_MAX_PROFILE_COLUMNS_SINGLE = 50
_MAX_PROFILE_ROWS = 10000
_MAX_PROFILE_TOP_VALUES = 50
_MAX_PROFILE_LOW_CARD_VALUES = 200
_LOW_CARD_UNIQUE_RATIO = 0.5
_CATEGORICAL_UNIQUE_RATIO = 0.3
_MAX_PROFILE_VALUE_CHARS = 80
_MAX_FILTER_SHEETS = 20
_MAX_FILTER_COLUMNS = 50
_MAX_FILTER_ROWS = 50000
_MAX_FILTER_SAMPLE_ROWS = 50
_MAX_FILTER_EXPORT_ROWS = 5000
_MAX_FILTER_RULES = 8
_COLUMN_LETTER_RE = re.compile(r"^[A-Za-z]{1,3}$")
_EXCEL_READ_ACTION_ALIASES = {
    "read": "inspect",
    "open": "inspect",
    "load": "inspect",
    "preview": "inspect",
    "view": "inspect",
    "get": "inspect",
}
_EXCEL_PROFILE_ACTION_ALIASES = {
    "profile",
    "analyze",
    "stats",
    "statistics",
    "summarize",
    "summary",
    "aggregate",
}
_EXCEL_FILTER_ACTION_ALIASES = {
    "filter",
    "query",
    "where",
    "select",
    "search",
}
_FILTER_OPS = frozenset({
    "eq",
    "ne",
    "contains",
    "not_contains",
    "startswith",
    "endswith",
    "gt",
    "gte",
    "lt",
    "lte",
    "in",
    "not_in",
    "empty",
    "not_empty",
})
_FILTER_OP_ALIASES = {
    "=": "eq",
    "==": "eq",
    "等于": "eq",
    "!=": "ne",
    "<>": "ne",
    "不等于": "ne",
    "包含": "contains",
    "like": "contains",
    "includes": "contains",
    "不包含": "not_contains",
    "startswith": "startswith",
    "starts_with": "startswith",
    "开头": "startswith",
    "endswith": "endswith",
    "ends_with": "endswith",
    "结尾": "endswith",
    ">": "gt",
    ">=": "gte",
    "<": "lt",
    "<=": "lte",
    "in": "in",
    "属于": "in",
    "not_in": "not_in",
    "不属于": "not_in",
    "empty": "empty",
    "blank": "empty",
    "空": "empty",
    "not_empty": "not_empty",
    "非空": "not_empty",
}


def _context():
    context = get_current_agent_context()
    if not context:
        raise DocumentPathError("无法获取当前执行上下文")
    return context


def _context_user_name(context) -> str | None:
    dims = context.user_dimensions or {}
    raw_name = dims.get("user_name") or dims.get("username")
    if not raw_name:
        return None
    name = str(raw_name).strip()
    return name or None


def _workspace_user_id(context) -> str:
    from app.services.ai.conversation_identity import session_user_id_from_agent_context

    return session_user_id_from_agent_context(context)


async def _input_path(path: str):
    context = _context()
    return await resolve_document_input_path(
        path,
        allowed_attachment_paths=context.authorized_attachment_paths,
        user_id=_workspace_user_id(context),
        conversation_id=context.conversation_id,
        allowed_extensions=_EXTENSIONS,
        user_name=_context_user_name(context),
    )


async def _output_path(filename: str):
    context = _context()
    return await resolve_document_output_path(
        filename,
        user_id=_workspace_user_id(context),
        conversation_id=context.conversation_id,
        allowed_extensions=_EXTENSIONS,
        user_name=_context_user_name(context),
    )


async def _artifact_result(output_path, *, summary: str, changes: dict[str, Any]) -> dict[str, Any]:
    context = _context()
    artifact = await register_artifact(
        source_path=output_path,
        filename=output_path.name,
        owner_user_id=context.user_id,
        artifact_type="excel",
        conversation_id=context.conversation_id,
        trace_id=context.trace_id,
    )
    return {
        "status": "ok",
        "summary": summary,
        "changes": changes,
        "artifact": artifact.to_tool_payload(),
    }


def _has_filter_payload(filters: Any) -> bool:
    if filters is None:
        return False
    if isinstance(filters, str):
        return bool(filters.strip())
    if isinstance(filters, dict):
        return bool(filters)
    if isinstance(filters, (list, tuple)):
        return len(filters) > 0
    return True


def _normalize_excel_read_action(
    action: str | None,
    sheet_name: str | None,
    cell_range: str | None,
    filters: Any = None,
) -> str:
    normalized = str(action or "").strip().lower()
    if normalized == "read_range":
        return "read_range"
    if normalized == "profile" or normalized in _EXCEL_PROFILE_ACTION_ALIASES:
        return "profile"
    if (
        normalized == "filter"
        or normalized in _EXCEL_FILTER_ACTION_ALIASES
        or _has_filter_payload(filters)
    ):
        return "filter"
    if sheet_name and cell_range:
        return "read_range"
    if not normalized or normalized == "inspect" or normalized in _EXCEL_READ_ACTION_ALIASES:
        return "inspect"
    raise DocumentPathError("excel_document_read 仅支持 inspect、profile、filter 或 read_range")


def _format_sheet_names(sheetnames: list[str]) -> str:
    names = [str(name) for name in sheetnames if str(name).strip()]
    if not names:
        return "（无）"
    return "、".join(names)


def _resolve_sheet_name(requested: str | None, sheetnames: list[str]) -> str:
    """Exact match, then unique prefix/contains. Ambiguous or missing names list all sheets."""
    wanted = str(requested or "").strip()
    if not wanted:
        raise DocumentPathError(
            f"需要工作表名称。当前工作簿的工作表：{_format_sheet_names(sheetnames)}"
        )
    names = [str(name) for name in sheetnames]
    if wanted in names:
        return wanted

    folded = wanted.casefold()
    case_hits = [name for name in names if name.casefold() == folded]
    if len(case_hits) == 1:
        return case_hits[0]
    if len(case_hits) > 1:
        raise DocumentPathError(
            f"工作表「{wanted}」匹配到多个同名表：{_format_sheet_names(case_hits)}"
        )

    prefix_hits = [name for name in names if name.startswith(wanted)]
    if len(prefix_hits) == 1:
        return prefix_hits[0]

    contains_hits = [name for name in names if wanted in name]
    if len(contains_hits) == 1:
        return contains_hits[0]

    if prefix_hits or contains_hits:
        candidates = prefix_hits or contains_hits
        raise DocumentPathError(
            f"工作表「{wanted}」不唯一，候选：{_format_sheet_names(candidates)}。"
            f"请用 inspect 或 profile 返回的完整表名。当前工作簿的工作表：{_format_sheet_names(names)}"
        )
    raise DocumentPathError(
        f"工作表「{wanted}」不存在。当前工作簿的工作表：{_format_sheet_names(names)}"
    )


def _cell_display(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, dt_time):
        return value.isoformat(timespec="seconds")
    return str(value).strip()


def _is_title_row(values: Any) -> bool:
    """A leading banner row is typically one long filled cell with the rest empty."""
    texts = [_cell_display(value) for value in values]
    nonempty = [text for text in texts if text]
    return len(nonempty) == 1 and len(nonempty[0]) >= 10


def _truncate_profile_value(value: str) -> str:
    if len(value) <= _MAX_PROFILE_VALUE_CHARS:
        return value
    return value[: _MAX_PROFILE_VALUE_CHARS - 3] + "..."


def _top_values_limit(unique_count: int, non_empty: int) -> int:
    """Low-cardinality categories return up to 200; high-cardinality stays at 50."""
    if unique_count <= 0:
        return 0
    if non_empty <= 0:
        return min(unique_count, _MAX_PROFILE_TOP_VALUES)
    ratio = unique_count / float(non_empty)
    if unique_count <= _MAX_PROFILE_LOW_CARD_VALUES and ratio < _LOW_CARD_UNIQUE_RATIO:
        return unique_count
    if ratio < _CATEGORICAL_UNIQUE_RATIO:
        return min(unique_count, _MAX_PROFILE_LOW_CARD_VALUES)
    return min(unique_count, _MAX_PROFILE_TOP_VALUES)


def _value_distribution(
    counter: Counter[str],
    non_empty: int | None = None,
) -> dict[str, Any]:
    unique_count = len(counter)
    filled = int(non_empty) if non_empty is not None else int(sum(counter.values()))
    limit = _top_values_limit(unique_count, filled)
    top = counter.most_common(limit) if limit else []
    covered_rows = sum(count for _value, count in top)
    omitted_unique = max(0, unique_count - len(top))
    coverage = round(covered_rows / filled, 4) if filled else 0.0
    return {
        "unique_count": unique_count,
        "top_values": [
            {"value": _truncate_profile_value(value), "count": count}
            for value, count in top
        ],
        "top_values_limit": limit,
        "omitted_unique": omitted_unique,
        "top_values_row_coverage": coverage,
        "top_values_complete": omitted_unique == 0,
    }


def _try_number(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        number = float(value)
        if number != number:
            return None
        return number
    return None


def _profile_bounds(
    worksheet,
    cell_range: str | None,
    *,
    max_columns: int,
) -> tuple[int, int, int, int, bool]:
    if cell_range:
        min_col, min_row, max_col, max_row = range_boundaries(cell_range)
    else:
        min_col = 1
        min_row = 1
        max_col = max(int(worksheet.max_column or 1), 1)
        max_row = max(int(worksheet.max_row or 1), 1)
    truncated = False
    if max_col - min_col + 1 > max_columns:
        max_col = min_col + max_columns - 1
        truncated = True
    if max_row - min_row + 1 > _MAX_PROFILE_ROWS:
        max_row = min_row + _MAX_PROFILE_ROWS - 1
        truncated = True
    return min_col, min_row, max_col, max_row, truncated


def _profile_worksheet(
    worksheet,
    *,
    min_row: int,
    max_row: int,
    min_col: int,
    max_col: int,
    truncated: bool,
) -> dict[str, Any]:
    counters: dict[int, Counter[str]] = {
        col: Counter() for col in range(min_col, max_col + 1)
    }
    empty_counts = {col: 0 for col in range(min_col, max_col + 1)}
    non_empty_counts = {col: 0 for col in range(min_col, max_col + 1)}
    numeric_min: dict[int, float] = {}
    numeric_max: dict[int, float] = {}
    numeric_sum: dict[int, float] = {}
    numeric_n: dict[int, int] = {}
    headers: dict[int, str] = {}
    scanned_rows = 0
    header_taken = False

    for row in worksheet.iter_rows(
        min_row=min_row,
        max_row=max_row,
        min_col=min_col,
        max_col=max_col,
        values_only=True,
    ):
        scanned_rows += 1
        if not header_taken:
            if _is_title_row(row):
                continue
            for offset, value in enumerate(row):
                col = min_col + offset
                label = _cell_display(value)
                headers[col] = label or get_column_letter(col)
            header_taken = True
            continue
        for offset, value in enumerate(row):
            col = min_col + offset
            text = _cell_display(value)
            if not text:
                empty_counts[col] += 1
                continue
            non_empty_counts[col] += 1
            counters[col][text] += 1
            number = _try_number(value)
            if number is None:
                continue
            numeric_n[col] = numeric_n.get(col, 0) + 1
            numeric_sum[col] = numeric_sum.get(col, 0.0) + number
            numeric_min[col] = (
                number if col not in numeric_min else min(numeric_min[col], number)
            )
            numeric_max[col] = (
                number if col not in numeric_max else max(numeric_max[col], number)
            )

    columns: list[dict[str, Any]] = []
    for col in range(min_col, max_col + 1):
        counter = counters[col]
        payload: dict[str, Any] = {
            "letter": get_column_letter(col),
            "header": headers.get(col) or get_column_letter(col),
            "non_empty": non_empty_counts[col],
            "empty": empty_counts[col],
            **_value_distribution(counter, non_empty_counts[col]),
        }
        numeric_count = numeric_n.get(col, 0)
        if numeric_count and numeric_count >= max(1, int(non_empty_counts[col] * 0.5)):
            payload["numeric"] = {
                "count": numeric_count,
                "min": numeric_min[col],
                "max": numeric_max[col],
                "avg": round(numeric_sum[col] / numeric_count, 4),
            }
        columns.append(payload)

    last_row = min_row + scanned_rows - 1 if scanned_rows else min_row
    return {
        "name": worksheet.title,
        "rows": worksheet.max_row,
        "columns": worksheet.max_column,
        "data_rows": max(0, scanned_rows - (1 if header_taken else 0)),
        "scanned_range": (
            f"{get_column_letter(min_col)}{min_row}:"
            f"{get_column_letter(max_col)}{last_row}"
        ),
        "truncated": truncated,
        "column_profiles": columns,
    }


def _try_number_loose(value: Any) -> float | None:
    number = _try_number(value)
    if number is not None:
        return number
    text = _cell_display(value).replace(",", "").replace("%", "").strip()
    if not text:
        return None
    try:
        parsed = float(text)
    except ValueError:
        return None
    if parsed != parsed:
        return None
    return parsed


def _try_datetime(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    text = _cell_display(value)
    if not text:
        return None
    for fmt in (
        "%Y-%m-%d",
        "%Y/%m/%d",
        "%Y.%m.%d",
        "%Y-%m-%d %H:%M:%S",
        "%Y/%m/%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%Y/%m/%d %H:%M",
    ):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def _normalize_filter_op(raw: Any) -> str:
    op = str(raw or "").strip().lower()
    if not op:
        return "eq"
    mapped = _FILTER_OP_ALIASES.get(op, op)
    if mapped not in _FILTER_OPS:
        allowed = "、".join(sorted(_FILTER_OPS))
        raise DocumentPathError(f"不支持的筛选运算符「{raw}」。可用：{allowed}")
    return mapped


def _split_in_values(value: Any) -> list[str]:
    if isinstance(value, (list, tuple, set)):
        return [str(item).strip() for item in value if str(item).strip()]
    text = str(value or "").strip()
    if not text:
        return []
    if "," in text:
        return [part.strip() for part in text.split(",") if part.strip()]
    if "，" in text:
        return [part.strip() for part in text.split("，") if part.strip()]
    return [text]


def _parse_filter_rules(raw: Any) -> list[dict[str, Any]]:
    payload = raw
    if isinstance(raw, str):
        text = raw.strip()
        if not text:
            payload = []
        else:
            try:
                payload = json.loads(text)
            except json.JSONDecodeError as exc:
                raise DocumentPathError(
                    "filters 必须是对象数组，例如 "
                    '[{"header":"列名","op":"contains","value":"关键词"}]'
                ) from exc
    if isinstance(payload, dict):
        payload = [payload]
    if not isinstance(payload, list) or not payload:
        raise DocumentPathError(
            "filter 需要 filters。每条规则包含 header 或 column，以及 op、value。"
            '示例：[{"header":"列名","op":"contains","value":"关键词"}]'
        )
    if len(payload) > _MAX_FILTER_RULES:
        raise DocumentPathError(f"筛选条件最多 {_MAX_FILTER_RULES} 条")
    rules: list[dict[str, Any]] = []
    for item in payload:
        if not isinstance(item, dict):
            raise DocumentPathError("filters 中的每一项必须是对象")
        column = str(
            item.get("column")
            or item.get("col")
            or item.get("header")
            or item.get("field")
            or item.get("name")
            or ""
        ).strip()
        op = _normalize_filter_op(item.get("op") or item.get("operator"))
        if not column:
            raise DocumentPathError("筛选规则需要 header 或 column")
        rules.append(
            {
                "column": column,
                "op": op,
                "value": item.get("value"),
            }
        )
    return rules


def _looks_like_column_letter(text: str) -> bool:
    if not _COLUMN_LETTER_RE.fullmatch(text or ""):
        return False
    try:
        return column_index_from_string(text.upper()) >= 1
    except ValueError:
        return False


def _resolve_filter_column(
    requested: str,
    headers: dict[int, str],
    *,
    min_col: int,
    max_col: int,
) -> int:
    wanted = str(requested or "").strip()
    if not wanted:
        raise DocumentPathError("筛选规则缺少列名")
    if _looks_like_column_letter(wanted):
        col = column_index_from_string(wanted.upper())
        if min_col <= col <= max_col:
            return col
        raise DocumentPathError(
            f"列 {wanted.upper()} 不在当前读取范围 "
            f"{get_column_letter(min_col)}:{get_column_letter(max_col)}"
        )
    names = [headers.get(col, get_column_letter(col)) for col in range(min_col, max_col + 1)]
    mapping = {
        headers.get(col, get_column_letter(col)): col
        for col in range(min_col, max_col + 1)
    }
    if wanted in mapping:
        return mapping[wanted]
    folded = wanted.casefold()
    case_hits = [name for name in names if name.casefold() == folded]
    if len(case_hits) == 1:
        return mapping[case_hits[0]]
    prefix_hits = [name for name in names if name.startswith(wanted)]
    if len(prefix_hits) == 1:
        return mapping[prefix_hits[0]]
    contains_hits = [name for name in names if wanted in name]
    if len(contains_hits) == 1:
        return mapping[contains_hits[0]]
    raise DocumentPathError(
        f"找不到列「{wanted}」。当前表头：{_format_sheet_names(names)}"
    )


def _cell_matches_rule(value: Any, op: str, expected: Any) -> bool:
    text = _cell_display(value)
    if op == "empty":
        return not text
    if op == "not_empty":
        return bool(text)
    expected_text = _cell_display(expected)
    if op == "eq":
        return text.casefold() == expected_text.casefold()
    if op == "ne":
        return text.casefold() != expected_text.casefold()
    if op == "contains":
        return bool(expected_text) and expected_text.casefold() in text.casefold()
    if op == "not_contains":
        return bool(expected_text) and expected_text.casefold() not in text.casefold()
    if op == "startswith":
        return bool(expected_text) and text.casefold().startswith(expected_text.casefold())
    if op == "endswith":
        return bool(expected_text) and text.casefold().endswith(expected_text.casefold())
    if op in {"in", "not_in"}:
        options = {item.casefold() for item in _split_in_values(expected)}
        hit = text.casefold() in options
        return hit if op == "in" else not hit
    left_dt = _try_datetime(value)
    right_dt = _try_datetime(expected)
    if left_dt is not None and right_dt is not None:
        pairs = {
            "gt": left_dt > right_dt,
            "gte": left_dt >= right_dt,
            "lt": left_dt < right_dt,
            "lte": left_dt <= right_dt,
        }
        return pairs[op]
    left_num = _try_number_loose(value)
    right_num = _try_number_loose(expected)
    if left_num is None or right_num is None:
        return False
    pairs = {
        "gt": left_num > right_num,
        "gte": left_num >= right_num,
        "lt": left_num < right_num,
        "lte": left_num <= right_num,
    }
    return pairs[op]


def _row_matches_rules(
    values_by_col: dict[int, Any],
    rules: list[dict[str, Any]],
    column_indexes: list[int],
    combine: str,
) -> bool:
    hits = []
    for rule, col in zip(rules, column_indexes):
        hits.append(_cell_matches_rule(values_by_col.get(col), rule["op"], rule["value"]))
    if combine == "or":
        return any(hits)
    return all(hits)


def _normalize_combine(raw: Any) -> str:
    text = str(raw or "and").strip().lower()
    if text in {"or", "any", "或"}:
        return "or"
    return "and"


def _excel_filter_bounds(
    worksheet,
    cell_range: str | None,
) -> tuple[int, int, int, int, bool]:
    if cell_range:
        min_col, min_row, max_col, max_row = range_boundaries(cell_range)
    else:
        min_col = 1
        min_row = 1
        max_col = max(int(worksheet.max_column or 1), 1)
        max_row = max(int(worksheet.max_row or 1), 1)
    truncated = False
    if max_col - min_col + 1 > _MAX_FILTER_COLUMNS:
        max_col = min_col + _MAX_FILTER_COLUMNS - 1
        truncated = True
    if max_row - min_row + 1 > _MAX_FILTER_ROWS:
        max_row = min_row + _MAX_FILTER_ROWS - 1
        truncated = True
    return min_col, min_row, max_col, max_row, truncated


def _sanitize_export_sheet_title(name: str) -> str:
    cleaned = re.sub(r"[\\/*?:\[\]]", "_", str(name or "").strip()) or "筛选结果"
    return cleaned[:31]


def _filter_worksheet(
    worksheet,
    *,
    rules: list[dict[str, Any]],
    combine: str,
    min_row: int,
    max_row: int,
    min_col: int,
    max_col: int,
    truncated: bool,
    collect_export: bool,
) -> dict[str, Any]:
    headers: dict[int, str] = {}
    column_indexes: list[int] | None = None
    matched_count = 0
    scanned_rows = 0
    header_taken = False
    sample_rows: list[list[str]] = []
    export_rows: list[list[Any]] = []
    header_values: list[str] = []
    matched_value_counts: dict[int, Counter[str]] = {}

    for row in worksheet.iter_rows(
        min_row=min_row,
        max_row=max_row,
        min_col=min_col,
        max_col=max_col,
        values_only=True,
    ):
        scanned_rows += 1
        if not header_taken:
            if _is_title_row(row):
                continue
            for offset, value in enumerate(row):
                col = min_col + offset
                label = _cell_display(value)
                headers[col] = label or get_column_letter(col)
            header_values = [
                headers[col] for col in range(min_col, max_col + 1)
            ]
            column_indexes = [
                _resolve_filter_column(
                    rule["column"],
                    headers,
                    min_col=min_col,
                    max_col=max_col,
                )
                for rule in rules
            ]
            matched_value_counts = {col: Counter() for col in range(min_col, max_col + 1)}
            header_taken = True
            continue
        values_by_col = {
            min_col + offset: value for offset, value in enumerate(row)
        }
        if not _row_matches_rules(values_by_col, rules, column_indexes or [], combine):
            continue
        matched_count += 1
        for col in range(min_col, max_col + 1):
            text = _cell_display(values_by_col.get(col))
            if text:
                matched_value_counts[col][text] += 1
        display_row = [_cell_display(value) for value in row]
        if len(sample_rows) < _MAX_FILTER_SAMPLE_ROWS:
            sample_rows.append(display_row)
        if collect_export and len(export_rows) < _MAX_FILTER_EXPORT_ROWS:
            export_rows.append(list(row))

    last_row = min_row + scanned_rows - 1 if scanned_rows else min_row
    return {
        "name": worksheet.title,
        "rows": worksheet.max_row,
        "columns": worksheet.max_column,
        "scanned_rows": max(0, scanned_rows - (1 if header_taken else 0)),
        "scanned_range": (
            f"{get_column_letter(min_col)}{min_row}:"
            f"{get_column_letter(max_col)}{last_row}"
        ),
        "matched_count": matched_count,
        "headers": header_values,
        "matched_values": [
            {
                "letter": get_column_letter(col),
                "header": headers.get(col) or get_column_letter(col),
                **_value_distribution(counter),
            }
            for col, counter in matched_value_counts.items()
        ],
        "sample_rows": sample_rows,
        "sample_truncated": matched_count > len(sample_rows),
        "export_truncated": collect_export and matched_count > len(export_rows),
        "truncated": truncated
        or matched_count > len(sample_rows)
        or (collect_export and matched_count > len(export_rows)),
        "export_rows": export_rows,
    }


async def _export_filtered_workbook(
    sheets: list[dict[str, Any]],
    output_filename: str,
) -> dict[str, Any]:
    workbook = Workbook()
    default_sheet = workbook.active
    first = True
    exported_rows = 0
    used_titles: set[str] = set()
    for item in sheets:
        title = _sanitize_export_sheet_title(str(item.get("name") or "筛选结果"))
        base = title
        suffix_n = 2
        while title.casefold() in used_titles:
            suffix = f"_{suffix_n}"
            title = (base[: 31 - len(suffix)] + suffix)
            suffix_n += 1
        used_titles.add(title.casefold())
        if first:
            worksheet = default_sheet
            worksheet.title = title
            first = False
        else:
            worksheet = workbook.create_sheet(title)
        headers = list(item.get("headers") or [])
        if headers:
            worksheet.append(headers)
        for row in item.get("export_rows") or []:
            worksheet.append(list(row))
            exported_rows += 1
    output_path = await _output_path(output_filename)
    workbook.save(output_path)
    workbook.close()
    return await _artifact_result(
        output_path,
        summary="已导出筛选结果",
        changes={"exported_rows": exported_rows, "exported_sheets": len(sheets)},
    )


@tool
async def excel_document_read(
    path: str,
    action: str = "inspect",
    sheet_name: str | None = None,
    cell_range: str | None = None,
    filters: list[dict[str, Any]] | dict[str, Any] | str | None = None,
    combine: str = "and",
    output_filename: str | None = None,
) -> dict[str, Any]:
    """Inspect, profile, filter or read a bounded range from an uploaded workbook.

    action:
    - inspect: list sheets and a small preview. Use this first. Do not pass action=read.
    - profile: scan a sheet or the whole workbook and return per-column counts,
      unique_count and adaptive top_values (up to 200 for low-cardinality
      categories, 50 for high-cardinality). Incomplete lists include
      omitted_unique and top_values_row_coverage. sheet_name is optional.
    - filter: evaluate structured predicates in-process. Do not page rows into
      the model. filters is a list of {header|column, op, value}. combine is
      and/or. matched_count is the total matching rows, not one object's count.
      matched_values covers every column with the same adaptive top_values.
      Also returns up to 50 sample rows. Pass output_filename to export matches.
    - read_range: preview cells only, not for analysis. Requires sheet_name
      and cell_range such as A1:G50.
      Prefer the exact title from inspect/profile. A unique prefix or substring
      of an existing sheet name is also accepted. Each call is capped at 1000
      rows and 50 columns; oversized ranges are truncated and return next_range.
    """
    action = _normalize_excel_read_action(action, sheet_name, cell_range, filters)
    input_path = await _input_path(path)
    workbook = load_workbook(input_path, read_only=True, data_only=False)
    try:
        if action == "inspect":
            sheets = []
            for worksheet in workbook.worksheets:
                preview = []
                headers: list[str] = []
                skipped_title = False
                max_preview = min(worksheet.max_row or 1, _MAX_PREVIEW_ROWS + 2)
                max_col = min(worksheet.max_column or 1, _MAX_PREVIEW_COLUMNS)
                for row in worksheet.iter_rows(
                    max_row=max_preview,
                    max_col=max_col,
                    values_only=True,
                ):
                    values = list(row)
                    if not headers and not skipped_title and _is_title_row(values):
                        skipped_title = True
                        continue
                    if not headers:
                        headers = [
                            _cell_display(value) or get_column_letter(index)
                            for index, value in enumerate(values, start=1)
                        ]
                    preview.append(values)
                    if len(preview) >= _MAX_PREVIEW_ROWS:
                        break
                sheets.append({
                    "name": worksheet.title,
                    "rows": worksheet.max_row,
                    "columns": worksheet.max_column,
                    "headers": headers,
                    "preview": preview,
                })
            return {"status": "ok", "summary": f"工作簿包含 {len(sheets)} 个工作表", "data": {"sheets": sheets}, "truncated": False}
        if action == "profile":
            targets = list(workbook.worksheets)
            if sheet_name:
                resolved = _resolve_sheet_name(sheet_name, list(workbook.sheetnames))
                targets = [workbook[resolved]]
            truncated = len(targets) > _MAX_PROFILE_SHEETS
            targets = targets[:_MAX_PROFILE_SHEETS]
            max_columns = (
                _MAX_PROFILE_COLUMNS_SINGLE if sheet_name else _MAX_PROFILE_COLUMNS
            )
            sheets = []
            any_truncated = truncated
            for worksheet in targets:
                min_col, min_row, max_col, max_row, range_truncated = _profile_bounds(
                    worksheet,
                    cell_range if sheet_name else None,
                    max_columns=max_columns,
                )
                any_truncated = any_truncated or range_truncated
                sheets.append(
                    _profile_worksheet(
                        worksheet,
                        min_row=min_row,
                        max_row=max_row,
                        min_col=min_col,
                        max_col=max_col,
                        truncated=range_truncated,
                    )
                )
            names = "、".join(item["name"] for item in sheets)
            summary = f"已统计 {len(sheets)} 张工作表"
            if names:
                summary += f"：{names}"
            incomplete = []
            for item in sheets:
                for column in item.get("column_profiles") or []:
                    omitted = int(column.get("omitted_unique") or 0)
                    if omitted <= 0:
                        continue
                    incomplete.append(
                        f"{item['name']}.{column.get('header')}"
                        f"列出{len(column.get('top_values') or [])}项/共{column.get('unique_count')}个"
                    )
            if incomplete:
                summary += (
                    "。部分列未列出全部取值（"
                    + "；".join(incomplete[:6])
                    + "），未列出对象请用 filter 且专有全称用 eq"
                )
            if any_truncated:
                summary += "（部分工作表或列已截取）"
            return {
                "status": "ok",
                "summary": summary,
                "data": {"sheets": sheets},
                "truncated": any_truncated,
            }
        if action == "filter":
            rules = _parse_filter_rules(filters)
            combine_mode = _normalize_combine(combine)
            targets = list(workbook.worksheets)
            if sheet_name:
                resolved = _resolve_sheet_name(sheet_name, list(workbook.sheetnames))
                targets = [workbook[resolved]]
            truncated = len(targets) > _MAX_FILTER_SHEETS
            targets = targets[:_MAX_FILTER_SHEETS]
            collect_export = bool(str(output_filename or "").strip())
            sheets = []
            skipped: list[dict[str, str]] = []
            any_truncated = truncated
            total_matched = 0
            for worksheet in targets:
                min_col, min_row, max_col, max_row, range_truncated = _excel_filter_bounds(
                    worksheet,
                    cell_range if sheet_name else None,
                )
                try:
                    item = _filter_worksheet(
                        worksheet,
                        rules=rules,
                        combine=combine_mode,
                        min_row=min_row,
                        max_row=max_row,
                        min_col=min_col,
                        max_col=max_col,
                        truncated=range_truncated,
                        collect_export=collect_export,
                    )
                except DocumentPathError as exc:
                    if sheet_name:
                        raise
                    skipped.append({"name": worksheet.title, "reason": str(exc)})
                    continue
                any_truncated = any_truncated or bool(item.get("truncated"))
                total_matched += int(item.get("matched_count") or 0)
                sheets.append(item)
            if not sheets and skipped:
                raise DocumentPathError(
                    "没有工作表包含指定筛选列。"
                    + " ".join(f"{item['name']}：{item['reason']}" for item in skipped)
                )
            workbook.close()
            workbook = None
            artifact = None
            if collect_export:
                artifact = await _export_filtered_workbook(
                    sheets,
                    str(output_filename).strip(),
                )
            public_sheets = []
            for item in sheets:
                public_item = dict(item)
                public_item.pop("export_rows", None)
                public_sheets.append(public_item)
            names = "、".join(item["name"] for item in public_sheets)
            summary = f"已筛选 {len(public_sheets)} 张工作表，共命中 {total_matched} 行"
            if names:
                summary += f"：{names}"
            mixed = []
            incomplete = []
            for item in public_sheets:
                for value_col in item.get("matched_values") or []:
                    unique = int(value_col.get("unique_count") or 0)
                    omitted = int(value_col.get("omitted_unique") or 0)
                    label = f"{item['name']}.{value_col.get('header')}"
                    if omitted > 0:
                        incomplete.append(f"{label}未列出{omitted}个值")
                    elif unique > 1 and value_col.get("top_values_complete"):
                        mixed.append(f"{label}={unique}个不同值")
            if mixed or incomplete:
                summary += (
                    "。matched_count 是条件命中总行数，不是单一对象的条数；"
                    "分对象请看 matched_values"
                )
                if mixed:
                    summary += f"（{'；'.join(mixed[:4])}）"
                if incomplete:
                    summary += (
                        f"；名单不完整时未列出对象须 filter+eq（"
                        f"{'；'.join(incomplete[:4])}）"
                    )
            if any_truncated:
                summary += f"（样本最多 {_MAX_FILTER_SAMPLE_ROWS} 行；完整结果请传 output_filename 导出）"
            payload: dict[str, Any] = {
                "sheets": public_sheets,
                "matched_count": total_matched,
                "combine": combine_mode,
                "filters": rules,
            }
            if skipped:
                payload["skipped_sheets"] = skipped
            result: dict[str, Any] = {
                "status": "ok",
                "summary": summary,
                "data": payload,
                "truncated": any_truncated,
            }
            if artifact:
                result["artifact"] = artifact.get("artifact")
                result["changes"] = artifact.get("changes")
                result["summary"] = f"{summary}；已导出筛选结果"
            return result
        if not sheet_name or not cell_range:
            raise DocumentPathError("read_range 需要 sheet_name 和 cell_range")
        sheet_name = _resolve_sheet_name(sheet_name, list(workbook.sheetnames))
        min_col, min_row, max_col, max_row = range_boundaries(cell_range)
        requested_rows = max_row - min_row + 1
        requested_cols = max_col - min_col + 1
        end_row = max_row
        end_col = max_col
        truncated = False
        if requested_rows > _MAX_RANGE_ROWS:
            end_row = min_row + _MAX_RANGE_ROWS - 1
            truncated = True
        if requested_cols > _MAX_RANGE_COLUMNS:
            end_col = min_col + _MAX_RANGE_COLUMNS - 1
            truncated = True
        worksheet = workbook[sheet_name]
        values = [
            list(row)
            for row in worksheet.iter_rows(
                min_row=min_row,
                max_row=end_row,
                min_col=min_col,
                max_col=end_col,
                values_only=True,
            )
        ]
        read_range = (
            f"{get_column_letter(min_col)}{min_row}:"
            f"{get_column_letter(end_col)}{end_row}"
        )
        payload: dict[str, Any] = {
            "values": values,
            "requested_range": cell_range,
            "read_range": read_range,
        }
        summary = f"已读取 {sheet_name}!{read_range}"
        if truncated:
            next_row = end_row + 1
            if next_row <= max_row:
                next_end = min(max_row, next_row + _MAX_RANGE_ROWS - 1)
                payload["next_range"] = (
                    f"{get_column_letter(min_col)}{next_row}:"
                    f"{get_column_letter(end_col)}{next_end}"
                )
            payload["limit"] = {
                "max_rows": _MAX_RANGE_ROWS,
                "max_columns": _MAX_RANGE_COLUMNS,
            }
            summary += (
                f"（已截取，原范围 {cell_range} 超过 {_MAX_RANGE_ROWS} 行或 "
                f"{_MAX_RANGE_COLUMNS} 列；下一段用 read_range "
                f"{payload.get('next_range') or '继续分页'}）"
            )
        return {
            "status": "ok",
            "summary": summary,
            "data": payload,
            "truncated": truncated,
        }
    finally:
        if workbook is not None:
            workbook.close()


@tool
async def excel_document_write(
    action: str,
    output_filename: str,
    path: str | None = None,
    sheet_name: str | None = None,
    cells: list[dict[str, Any]] | None = None,
    rows: list[list[Any]] | None = None,
) -> dict[str, Any]:
    """Create or modify an Excel workbook copy and return a download link.

    For ``action="create"``, pass ``rows`` and/or ``cells`` to populate the
    new workbook's active sheet in the same call. A create call without either
    parameter intentionally produces a blank workbook.

    Copy artifact.download_url verbatim in the final response; do not alter its
    protocol, host, path, or token.
    """
    if action not in {"create", "write_cells", "append_rows", "create_sheet"}:
        raise DocumentPathError("excel_document_write 不支持该操作")
    if action == "create":
        workbook = Workbook()
        worksheet = workbook.active
        if sheet_name:
            worksheet.title = sheet_name
    else:
        if not path:
            raise DocumentPathError("修改工作簿需要 path")
        workbook = load_workbook(await _input_path(path), data_only=False)
        if not sheet_name:
            raise DocumentPathError("修改工作簿需要 sheet_name")
        if action != "create_sheet":
            sheet_name = _resolve_sheet_name(sheet_name, list(workbook.sheetnames))
    changes: dict[str, Any] = {}
    if action == "create":
        if cells:
            for cell in cells:
                address = str(cell.get("address") or "")
                if not address:
                    raise DocumentPathError("单元格地址不能为空")
                worksheet[address] = cell.get("value")
            changes["written_cells"] = len(cells)
        if rows:
            for row in rows:
                worksheet.append(list(row))
            changes["appended_rows"] = len(rows)
        changes["created_workbook"] = True
    elif action == "write_cells":
        worksheet = workbook[sheet_name]
        if not cells:
            raise DocumentPathError("write_cells 需要 cells")
        for cell in cells:
            address = str(cell.get("address") or "")
            if not address:
                raise DocumentPathError("单元格地址不能为空")
            worksheet[address] = cell.get("value")
        changes["written_cells"] = len(cells)
    elif action == "append_rows":
        worksheet = workbook[sheet_name]
        if not rows:
            raise DocumentPathError("append_rows 需要 rows")
        for row in rows:
            worksheet.append(list(row))
        changes["appended_rows"] = len(rows)
    elif action == "create_sheet":
        if sheet_name in workbook.sheetnames:
            raise DocumentPathError("工作表已存在")
        workbook.create_sheet(sheet_name)
        changes["created_sheet"] = sheet_name
    output_path = await _output_path(output_filename)
    workbook.save(output_path)
    workbook.close()
    return await _artifact_result(output_path, summary="已生成 Excel 文件", changes=changes)
